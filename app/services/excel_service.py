from openpyxl import Workbook
from openpyxl.styles import Font


def export_payments_excel(payments, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = "Payments"

    headers = [
        "Payment ID",
        "Farmer ID",
        "From Date",
        "To Date",
        "Total Liters",
        "Total Amount",
        "Status"
    ]

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)

    row = 2

    for payment in payments:
        ws.cell(row=row, column=1).value = payment.id
        ws.cell(row=row, column=2).value = payment.farmer_id
        ws.cell(row=row, column=3).value = str(payment.from_date)
        ws.cell(row=row, column=4).value = str(payment.to_date)
        ws.cell(row=row, column=5).value = payment.total_liters
        ws.cell(row=row, column=6).value = payment.total_amount
        ws.cell(row=row, column=7).value = payment.status
        row += 1

    # Auto-size columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = max_length + 3

    wb.save(filename)